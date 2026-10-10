"""European directory adapters with explicitly configured publishers.

Parsing never activates a dataset."""

import csv
import io
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from bs4 import BeautifulSoup
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xlrd import XLRDError, open_workbook

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100_000
BIC = re.compile(r"[A-Z0-9]{4}[A-Z]{2}[A-Z0-9]{2}(?:[A-Z0-9]{3})?")


@dataclass(frozen=True)
class Source:
    publisher: str
    page: str
    download: str | None
    hosts: tuple[str, ...]
    format: str
    code_pattern: str


SOURCES = {
    "NL": Source(
        "Betaalvereniging Nederland (provisional industry source)",
        "https://www.betaalvereniging.nl/kennisbank/iban-en-bic/",
        "https://www.betaalvereniging.nl/wp-content/uploads/2025/11/BIC-lijst-NL.xlsx",
        ("www.betaalvereniging.nl",),
        "xlsx",
        r"[A-Z]{4}",
    ),
    "CH": Source(
        "SIX",
        "https://www.six-group.com/de/products-services/banking-services/"
        "interbank-clearing/online-services/download-bank-master.html",
        "https://api.six-group.com/api/epcd/bankmaster/v3/bankmaster_V3.csv",
        ("www.six-group.com", "api.six-group.com"),
        "csv",
        r"[0-9]{5}",
    ),
    "PL": Source(
        "Narodowy Bank Polski",
        "https://ewib.nbp.pl/faces/pages/faq.xhtml",
        "https://ewib.nbp.pl/api/v1/zapytanie1/?format=json",
        ("ewib.nbp.pl",),
        "json",
        r"[0-9]{8}",
    ),
    "LT": Source(
        "Bank of Lithuania",
        "https://www.lb.lt/en/iban-and-financial-institution-codes",
        None,
        ("www.lb.lt",),
        "csv",
        r"[0-9]{5}",
    ),
    "BE": Source(
        "National Bank of Belgium",
        "https://www.nbb.be/en/payments-and-securities/bank-identification-codes",
        "https://www.nbb.be/doc/be/be/protocol/full_list_current.xlsx",
        ("www.nbb.be",),
        "xlsx",
        r"[0-9]{3}",
    ),
    "CZ": Source(
        "Czech National Bank",
        "https://www.cnb.cz/cs/platebni-styk/ucty-kody-bank/",
        "https://www.cnb.cz/cs/platebni-styk/.galleries/ucty_kody_bank/download/kody_bank_CR.csv",
        ("www.cnb.cz",),
        "csv",
        r"[0-9]{4}",
    ),
    "LV": Source(
        "Latvijas Banka",
        "https://www.bank.lv/en/operational-areas/payment-systems/identification-of-bic-by-iban",
        None,
        ("www.bank.lv",),
        "xls",
        r"[A-Z]{4}",
    ),
    "SI": Source(
        "Banka Slovenije",
        "https://www.bsi.si/sl/placilni-sistemi/placilni-in-transakcijski-racun",
        "https://www.bsi.si/sl/placilni-sistemi/placilni-in-transakcijski-racun",
        ("www.bsi.si",),
        "html",
        r"(?:[0-9]{2}|[0-9]{5})",
    ),
    "GR": Source(
        "Bank of Greece",
        "https://www.bankofgreece.gr/en/main-tasks/payment-systems-and-settlements/sepa",
        "https://www.bankofgreece.gr/RelatedDocuments/en-BIC-from-IBAN.xlsx",
        ("www.bankofgreece.gr",),
        "xlsx",
        r"[0-9]{3}",
    ),
}


@dataclass(frozen=True)
class Record:
    bank_code: str
    name: str
    bic: str | None
    postal_code: str = ""
    city: str = ""
    change_flag: str = "U"
    deletion_announced: bool = False
    successor_bank_code: str | None = None


@dataclass(frozen=True)
class Directory:
    records: list[Record]
    total_rows: int
    ignored_rows: int = 0
    published_on: date | None = None


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if int(value) != value:
            raise ValueError("Non-integer identifier in directory.")
        return str(int(value))
    if not isinstance(value, str):
        raise ValueError("Unexpected field type in directory.")
    return value.strip()


def bic(value) -> str | None:
    value = "".join(text(value).split())
    if value in {"", "/", "N/A", "nav", "NAV", "NYA", "-"}:
        return None
    if not BIC.fullmatch(value):
        raise ValueError("Invalid BIC in directory.")
    return value


def numeric(value, width: int) -> str:
    value = text(value)
    if not re.fullmatch(r"[0-9]{1," + str(width) + "}", value):
        raise ValueError("Invalid numeric bank identifier.")
    return value.zfill(width)


def csv_rows(content: bytes, encoding: str = "utf-8-sig") -> list[list[str]]:
    try:
        rows = list(csv.reader(io.StringIO(content.decode(encoding)), delimiter=";", strict=True))
    except (UnicodeError, csv.Error) as exc:
        raise ValueError("Invalid CSV file or character encoding.") from exc
    if len(rows) < 2 or any(not row for row in rows):
        raise ValueError("Empty rows or missing data in CSV.")
    return rows


def spreadsheet_rows(content: bytes, legacy: bool = False) -> list[list]:
    try:
        if legacy:
            book = open_workbook(file_contents=content)
            sheets = [s for s in book.sheets() if s.nrows]
            if len(sheets) != 1 or sheets[0].nrows > MAX_ROWS:
                raise ValueError("Exactly one non-empty worksheet is required.")
            return [sheets[0].row_values(i) for i in range(sheets[0].nrows)]
        # Bound decompression before openpyxl parses XML (zip bombs / huge sheets).
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_BYTES * 5:
                raise ValueError("Uncompressed XLSX file is too large.")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            sheets = [s for s in book if s.max_row and s.max_column]
            if len(sheets) != 1 or sheets[0].max_row > MAX_ROWS or sheets[0].max_column > 40:
                raise ValueError("Unexpected table structure.")
            rows = []
            for cells in sheets[0].iter_rows():
                if any(c.data_type in {"f", "e"} for c in cells):
                    raise ValueError("Formula or error cell in source file.")
                rows.append([c.value for c in cells])
            return rows
        finally:
            book.close()
    except (
        BadZipFile,
        XLRDError,
        KeyError,
        OSError,
        ParseError,
        InvalidFileException,
        DefusedXmlException,
    ) as exc:
        raise ValueError("Invalid spreadsheet file.") from exc


def parse_ch(content: bytes) -> Directory:
    rows = csv_rows(content)
    headers = [
        "IID/QR-IID",
        "Valid on",
        "Concatenation",
        "New IID/QR-IID",
        "SIC IID",
        "Headquarters",
        "IID type",
        "QR-IID allocation",
        "Name of bank/institution",
        "Street Name",
        "Building Number",
        "Post Code",
        "Town Name",
        "Country",
        "BIC",
        "SIC participation",
        "RTGS customer payments, CHF",
        "IP customer payments, CHF",
        "euroSIC participation",
        "LSV+/BDD, CHF",
        "LSV+/BDD, EUR",
    ]
    if rows[0][:-1] != headers or not re.fullmatch(r"[0-9]{14,15}", rows[0][-1]):
        raise ValueError("Unknown Swiss Bank Master v3 header.")
    records, dates = [], set()
    for r in rows[1:]:
        if len(r) != 21:
            raise ValueError("Swiss Bank Master: incorrect column count.")
        code = numeric(r[0], 5)
        dates.add(date.fromisoformat(r[1]))
        if r[2] == "Y":
            if any(r[4:]):
                raise ValueError("Concatenated IID contains unexpected active fields.")
            records.append(
                Record(code, "", None, change_flag="D", successor_bank_code=numeric(r[3], 5))
            )
        elif r[2] == "N" and r[6] in {"1", "2", "4"}:
            numeric(r[4], 6)
            numeric(r[5], 5)
            if r[3] or any(v not in {"Y", "N"} for v in r[15:]):
                raise ValueError("Invalid SIX status fields.")
            if r[6] == "4":
                numeric(r[7], 5)
                if not 30000 <= int(code) <= 31999:
                    raise ValueError("QR-IID outside the allowed range.")
            records.append(Record(code, r[8], bic(r[14]), r[11], r[12]))
        else:
            raise ValueError("Unknown IID status.")
    if len(dates) != 1:
        raise ValueError("Inconsistent SIX validity date.")
    return Directory(records, len(rows) - 1, published_on=dates.pop())


def parse_pl(content: bytes) -> Directory:
    def unique_pairs(pairs):
        obj = {}
        for key, value in pairs:
            if key in obj:
                raise ValueError("Duplicate JSON key.")
            obj[key] = value
        return obj

    try:
        data = json.loads(content, object_pairs_hook=unique_pairs)
        owners = data["listaWlascicieli"]
        if not isinstance(owners, list):
            raise ValueError("EWIB: missing owner list.")
        records = []
        for owner in owners:
            name = text(owner["nazwa"])
            units = owner.get("jednostki", [])
            if owner.get("siedziba"):
                units = [owner["siedziba"], *units]
            for unit in units:
                for entry in unit.get("numeryRozliczeniowe", []):
                    bics = {
                        bic(b["numer"]) for b in entry.get("numeryBic", []) if b["nazwa"] == "BIC"
                    }
                    if len(bics) > 1:
                        raise ValueError("EWIB: ambiguous BIC.")
                    records.append(
                        Record(
                            numeric(entry["numer"], 8),
                            name,
                            next(iter(bics), None),
                            text(unit.get("kodPocztowy")),
                            text(unit.get("miejscowosc")),
                        )
                    )
        return Directory(records, len(records))
    except (KeyError, TypeError, UnicodeError, json.JSONDecodeError) as exc:
        raise ValueError("Unknown EWIB JSON schema.") from exc


def parse_lt(content: bytes) -> Directory:
    # The English export inspected on 2026-10-08 is Windows-1257, not UTF-8.
    rows = csv_rows(content, "cp1257")
    if rows[0] != [
        "National ID",
        "BIC Code",
        "Financial Institution Name",
        "Branch Name",
        "Legal Entity Code",
        "City",
        "Branch Address",
        "Zip Code",
        "Location",
        "Country",
        "",
        "",
        "",
        "",
        "",
    ]:
        raise ValueError("Unknown Lithuanian CSV header (English export required).")
    records = []
    # This export ends with a fully empty padded row; internal gaps still fail.
    while len(rows) > 1 and not any(rows[-1]):
        rows.pop()
    for r in rows[1:]:
        if len(r) != 15 or any(r[10:]):
            raise ValueError("Unknown Lithuanian CSV columns.")
        records.append(Record(numeric(r[0], 5), r[2], bic(r[1]), r[7], r[5]))
    return Directory(records, len(records))


def parse_cz(content: bytes) -> Directory:
    rows = csv_rows(content)
    if rows[0] != [
        "Kód platebního styku",
        "Poskytovatel platebních služeb",
        "BIC kód (SWIFT)",
        "Systém CERTIS",
    ]:
        raise ValueError("Unknown CNB CSV header.")
    records = []
    for r in rows[1:]:
        if len(r) != 4 or r[3].strip() not in {"A", "-"}:
            raise ValueError("Invalid CNB row.")
        records.append(Record(numeric(r[0], 4), r[1], bic(r[2])))
    return Directory(records, len(records))


def parse_be(content: bytes) -> Directory:
    rows = spreadsheet_rows(content)
    if len(rows) < 3 or rows[1] != [
        "T_Identification_Number",
        "Biccode",
        "T_Institutions_Dutch",
        "T_Institutions_French",
        "T_Institutions_German",
        "T_Institutions_English",
    ]:
        raise ValueError("Unknown NBB XLSX header.")
    version = re.fullmatch(r"Version ([0-9]{2})/([0-9]{2})/([0-9]{4})", text(rows[0][0]))
    if not version:
        raise ValueError("Missing NBB version date.")
    published = date(int(version[3]), int(version[2]), int(version[1]))
    records, ignored = [], 0
    for r in rows[2:]:
        code = numeric(r[0], 3)
        names = [text(n) for n in r[2:]]
        # Unallocated and unavailable identifiers are not financial institutions.
        if names[0] in {"VRIJ", "Onbeschikbaar"}:
            ignored += 1
            continue
        name = names[3] or names[0] or names[1] or names[2]
        records.append(Record(code, name, bic(r[1])))
    return Directory(records, len(rows) - 2, ignored, published)


def parse_lv(content: bytes) -> Directory:
    rows = spreadsheet_rows(content, legacy=True)
    if len(rows) < 3 or rows[1] != ["", "Payment service provider", "IBAN structure", "BIC"]:
        raise ValueError("Unknown Latvijas Banka XLS header.")
    records = []
    for r in rows[2:]:
        structure = "".join(text(r[2]).split())
        if not re.fullmatch(r"LV\*\*[A-Z]{4}\*{13}", structure):
            raise ValueError("Invalid Latvian IBAN structure.")
        value = bic(r[3])
        code = structure[4:8]
        if not value or value[:4] != code or value[4:6] != "LV":
            raise ValueError("Conflicting IBAN/BIC mapping.")
        records.append(Record(code, text(r[1]), value))
    return Directory(records, len(records))


def parse_si(content: bytes) -> Directory:
    soup = BeautifulSoup(content.decode("utf-8-sig"), "html.parser")
    tables = [t for t in soup.find_all("table") if "BIC KODA" in t.get_text(" ", strip=True)]
    if len(tables) != 1:
        raise ValueError("Slovenian BIC table is missing or ambiguous.")
    records, spans = [], {}
    for row in tables[0].find_all("tr")[1:]:
        cells = iter(row.find_all(["td", "th"], recursive=False))
        values = []
        for col in range(4):
            if col in spans:
                value, remaining = spans[col]
                if remaining == 1:
                    del spans[col]
                else:
                    spans[col] = (value, remaining - 1)
            else:
                cell = next(cells, None)
                if cell is None or cell.get("colspan", "1") != "1":
                    raise ValueError("Unknown Slovenian table structure.")
                value = cell.get_text(" ", strip=True)
                count = int(cell.get("rowspan", "1"))
                if not 1 <= count <= 10:
                    raise ValueError("Invalid rowspan.")
                if count > 1:
                    spans[col] = (value, count - 1)
            values.append(value)
        if next(cells, None) is not None:
            raise ValueError("Unexpected extra Slovenian table column.")
        records.append(Record(values[3], values[0], bic(values[2])))
    if spans:
        raise ValueError("Truncated Slovenian table.")
    return Directory(records, len(records))


def parse_gr(content: bytes) -> Directory:
    # Header-based adapter; activation additionally requires real-source review.
    rows = spreadsheet_rows(content)
    header = None
    for index, row in enumerate(rows[:20]):
        labels = [re.sub(r"\s+", " ", text(v)).lower() for v in row]
        code = [i for i, v in enumerate(labels) if v in {"bank identifier", "bank identifiers"}]
        bics = [i for i, v in enumerate(labels) if v in {"bic", "bic code"}]
        names = [i for i, v in enumerate(labels) if v in {"bank", "bank name", "name", "psp"}]
        if len(code) == len(bics) == len(names) == 1:
            header = index, code[0], bics[0], names[0]
            break
    if header is None:
        raise ValueError("Unknown Bank of Greece XLSX header; source review required.")
    index, code, bics, names = header
    records = []
    for r in rows[index + 1 :]:
        if not any(text(v) for v in r):
            continue
        records.append(Record(numeric(r[code], 3), text(r[names]), bic(r[bics])))
    return Directory(records, len(records))


def parse_nl(content: bytes) -> Directory:
    rows = spreadsheet_rows(content)
    title = text(rows[0][0])
    match = re.fullmatch(
        r"BIC-lijst-NL \| BIC-list-NL \(Laatste update \| last update "
        r"([0-9]{2}-[0-9]{2}-[0-9]{4})\)",
        title,
    )
    if not match or any(text(v) for v in rows[0][1:]):
        raise ValueError("Unknown Dutch BIC list title or publication date.")

    published_on = datetime.strptime(match[1], "%d-%m-%Y").date()
    if [text(v) for v in rows[1]] != [
        "BIC",
        "Identifier",
        "Betaaldienstverlener / Payment Service Provider",
    ]:
        raise ValueError("Unknown Dutch BIC list header.")
    records = []
    for row in rows[2:]:
        if len(row) != 3 or not all(text(v) for v in row):
            raise ValueError("Incomplete Dutch BIC list row.")
        bank_bic, identifier, name = bic(row[0]), text(row[1]), text(row[2])
        if not re.fullmatch(r"[A-Z]{4}", identifier) or bank_bic[:4] != identifier:
            raise ValueError("Conflicting Dutch bank identifier and BIC.")
        if bank_bic[4:6] != "NL":
            raise ValueError("Dutch BIC must use country code NL.")
        records.append(Record(identifier, name, bank_bic))
    return Directory(records, len(records), published_on=published_on)


PARSERS = {
    "NL": parse_nl,
    "CH": parse_ch,
    "PL": parse_pl,
    "LT": parse_lt,
    "BE": parse_be,
    "CZ": parse_cz,
    "LV": parse_lv,
    "SI": parse_si,
    "GR": parse_gr,
}


def parse_directory(country: str, content: bytes) -> Directory:
    if country not in SOURCES or not content or len(content) > MAX_BYTES:
        raise ValueError("Unknown country, empty or oversized source file.")
    try:
        result = PARSERS[country](content)
    except (UnicodeError, IndexError, AttributeError) as exc:
        raise ValueError("Unknown or damaged source structure.") from exc
    seen = set()
    for record in result.records:
        if (
            not re.fullmatch(SOURCES[country].code_pattern, record.bank_code)
            or record.bank_code in seen
            or len(record.name) > 300
            or len(record.city) > 200
            or len(record.postal_code) > 32
            or (not record.name and record.change_flag != "D")
            or any(ord(c) < 32 for c in record.name + record.city + record.postal_code)
        ):
            raise ValueError("Invalid or duplicate bank record.")
        seen.add(record.bank_code)
    if not result.records or len(result.records) > MAX_ROWS:
        raise ValueError("No bank data or too many records.")
    return result
