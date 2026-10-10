import io
import json
from datetime import timedelta
from pathlib import Path
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select

from openiban.api import create_app
from openiban.directories import MAX_BYTES, SOURCES, parse_directory
from openiban.directory_download import OfficialRedirects, discover_url, validate_url
from openiban.directory_storage import (
    activate_directory,
    history,
    list_directory_versions,
    lookup_directory,
    stage_directory,
)
from openiban.storage import lookup, today
from openiban.validation import COUNTRIES, validate_iban


def xlsx(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


CH_HEADER = (
    "IID/QR-IID;Valid on;Concatenation;New IID/QR-IID;SIC IID;Headquarters;IID type;"
    "QR-IID allocation;Name of bank/institution;Street Name;Building Number;Post Code;Town Name;"
    "Country;BIC;SIC participation;RTGS customer payments, CHF;IP customer payments, CHF;"
    "euroSIC participation;LSV+/BDD, CHF;LSV+/BDD, EUR;202610071630000\r\n"
)
LT_HEADER = (
    "National ID;BIC Code;Financial Institution Name;Branch Name;Legal Entity Code;City;"
    "Branch Address;Zip Code;Location;Country;;;;;\r\n"
)
CZ_HEADER = "Kód platebního styku;Poskytovatel platebních služeb;BIC kód (SWIFT);Systém CERTIS\r\n"
BE_HEADER = [
    "T_Identification_Number",
    "Biccode",
    "T_Institutions_Dutch",
    "T_Institutions_French",
    "T_Institutions_German",
    "T_Institutions_English",
]


def content(country):
    if country == "NL":
        return xlsx(
            [
                [
                    "BIC-lijst-NL | BIC-list-NL (Laatste update | last update "
                    + today().strftime("%d-%m-%Y")
                    + ")"
                ],
                ["BIC", "Identifier", "Betaaldienstverlener / Payment Service Provider "],
                ["TESTNL22", "TEST", "Synthetic Bank"],
            ]
        )
    if country == "CH":
        return (
            CH_HEADER + f"100;{today()};N;;001008;100;1;;Synthetic Bank;Street;1;8000;"
            "Zürich;CH;TESTCH22XXX;Y;Y;Y;N;N;N\r\n"
        ).encode()
    if country == "PL":
        return json.dumps(
            {
                "listaWlascicieli": [
                    {
                        "nazwa": "Synthetic Bank Łódź",
                        "siedziba": {
                            "miejscowosc": "Łódź",
                            "kodPocztowy": "00-001",
                            "numeryRozliczeniowe": [
                                {
                                    "numer": "10100000",
                                    "numeryBic": [
                                        {"nazwa": "BIC", "numer": "TESTPL22XXX"},
                                        {"nazwa": "BIC SEPA", "numer": "ALTAPL22XXX"},
                                    ],
                                }
                            ],
                        },
                        "jednostki": [
                            {
                                "miejscowosc": "Warszawa",
                                "numeryRozliczeniowe": [{"numer": "10100001", "numeryBic": []}],
                            }
                        ],
                    }
                ]
            }
        ).encode()
    if country == "LT":
        return (
            LT_HEADER
            + "10100;TESTLT22;Synthetic Bank;;;Vilnius;;01001;;Lithuania;;;;;\r\n"
            + ";" * 14
            + "\r\n"
        ).encode("cp1257")
    if country == "BE":
        return xlsx(
            [
                ["Version " + today().strftime("%d/%m/%Y")],
                BE_HEADER,
                ["000", "TEST BE 22", "Synthetic Bank", "", "", ""],
                ["001", "VRIJ", "VRIJ", "LIBRE", "", ""],
                ["002", "N/A", "Onbeschikbaar", "Indisponible", "", ""],
                ["003", "NAV", "Another Bank", "", "", ""],
            ]
        )
    if country == "CZ":
        return (CZ_HEADER + "0100;Synthetic Bank;TESTCZ22;A\r\n").encode("utf-8-sig")
    if country == "LV":
        return Path(__file__).with_name("fixtures").joinpath("lv-synthetic.xls").read_bytes()
    if country == "SI":
        return b"""<html><table><tr><th>NAZIV</th><th>NASLOV</th><th>BIC KODA</th>
        <th>Dvo- oziroma petmestna identifikacijska oznaka</th></tr>
        <tr><td rowspan="2">Synthetic Bank</td><td rowspan="2">Address</td>
        <td>TESTSI22</td><td>01</td></tr><tr><td>TESTSI22</td><td>02</td></tr>
        <tr><td>Other Bank</td><td>Address</td><td>/</td><td>91001</td></tr></table></html>"""
    return xlsx([["Bank", "BIC", "Bank Identifier"], ["Synthetic Bank", "TESTGR22", "011"]])


@pytest.mark.parametrize("country", SOURCES)
def test_every_adapter_accepts_synthetic_native_format(country):
    parsed = parse_directory(country, content(country))
    assert parsed.records[0].name.startswith("Synthetic Bank")
    assert parsed.records[0].bic.startswith("TEST")


@pytest.mark.parametrize("country", SOURCES)
@pytest.mark.parametrize(
    "data", [b"", b"<html>Access denied</html>", b"garbage", b"x" * (MAX_BYTES + 1)]
)
def test_empty_error_page_and_oversized_sources_fail(country, data):
    with pytest.raises(ValueError):
        parse_directory(country, data)


def test_six_padding_qr_and_concatenated_iid():
    data = (
        content("CH")
        + (
            f"30000;{today()};N;;300005;30000;4;100;QR Bank;Street;1;"
            "8000;Zürich;CH;TESTCH22XXX;Y;Y;Y;N;N;N\r\n"
        ).encode()
    )
    data += ("30100;" + str(today()) + ";Y;30000;" + ";" * 16 + "\r\n").encode()
    parsed = parse_directory("CH", data)
    assert parsed.records[0].bank_code == "00100"
    assert parsed.records[1].bank_code == "30000"
    assert parsed.records[2].change_flag == "D"
    assert parsed.records[2].successor_bank_code == "30000"
    with pytest.raises(ValueError):
        parse_directory("CH", data.replace(b"30000;4;100", b"30000;9;100"))


def test_polish_branches_and_bic_kind():
    parsed = parse_directory("PL", content("PL"))
    assert [r.bank_code for r in parsed.records] == ["10100000", "10100001"]
    assert parsed.records[0].bic == "TESTPL22XXX"
    assert parsed.records[1].bic is None
    with pytest.raises(ValueError):
        parse_directory("PL", b'{"listaWlascicieli": [], "listaWlascicieli": []}')


def test_belgian_free_and_unavailable_codes_are_excluded():
    parsed = parse_directory("BE", content("BE"))
    assert [r.bank_code for r in parsed.records] == ["000", "003"]
    assert parsed.ignored_rows == 2
    assert parsed.records[1].bic is None


def test_slovenian_rowspans_and_optional_bic():
    parsed = parse_directory("SI", content("SI"))
    assert [r.bank_code for r in parsed.records] == ["01", "02", "91001"]
    assert parsed.records[1].name == parsed.records[0].name
    assert parsed.records[2].bic is None
    with pytest.raises(ValueError):
        parse_directory("SI", content("SI").replace(b'rowspan="2"', b'rowspan="3"'))


@pytest.mark.parametrize("country", ["CH", "CZ", "LT"])
def test_duplicate_codes_and_schema_drift_fail(country):
    data = content(country)
    line = data.splitlines()[1]
    with pytest.raises(ValueError):
        parse_directory(country, data + line + b"\r\n")
    with pytest.raises(ValueError):
        parse_directory(country, data.replace(b";", b",", 1))


def test_xlsx_formulas_and_zip_bombs_fail():
    with pytest.raises(ValueError):
        parse_directory(
            "GR", xlsx([["Bank", "BIC", "Bank Identifier"], ["=1+1", "TESTGR22", "011"]])
        )
    output = io.BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("large", b"0" * (MAX_BYTES * 5 + 1))
    with pytest.raises(ValueError):
        parse_directory("GR", output.getvalue())


@pytest.mark.parametrize("country", SOURCES)
def test_stage_activation_idempotency_and_country_isolation(engine, tmp_path, country, import_file):
    german = import_file(make_active=True)
    path = tmp_path / country
    path.write_bytes(content(country))
    url = SOURCES[country].download or SOURCES[country].page
    result = stage_directory(engine, country, path, today(), today() + timedelta(days=30), url)
    assert lookup_directory(engine, country, "00100")[0] is None
    repeated = stage_directory(engine, country, path, today(), today() + timedelta(days=30), url)
    assert repeated["version"] == result["version"]
    assert repeated["already_imported"]
    activate_directory(engine, country, result["version"], review_note="Synthetic source review")
    code = parse_directory(country, content(country)).records[0].bank_code
    assert lookup_directory(engine, country, code)[1]["name"].startswith("Synthetic Bank")
    assert lookup(engine, "12345678")[0]["version"] == german["version"]
    other = next(c for c in SOURCES if c != country)
    with pytest.raises(ValueError):
        activate_directory(engine, other, result["version"], review_note="Reviewed")
    path.write_bytes(b"BROKEN")
    with pytest.raises(ValueError):
        stage_directory(engine, country, path, today(), today(), url)
    assert len(list_directory_versions(engine, country)) == 1
    assert lookup_directory(engine, country, code)[0]["version"] == result["version"]


def test_future_expiry_review_and_rollback(engine, tmp_path):
    path = tmp_path / "cz.csv"
    path.write_bytes(content("CZ"))
    url = SOURCES["CZ"].download
    first = stage_directory(engine, "CZ", path, today(), today(), url)["version"]
    with pytest.raises(ValueError):
        activate_directory(engine, "CZ", first, review_note=" ")
    activate_directory(engine, "CZ", first, review_note="Source and use checked")
    assert lookup_directory(engine, "CZ", "0100", today() + timedelta(days=1))[1] is None
    with pytest.raises(ValueError):
        activate_directory(
            engine, "CZ", first, review_note="Rollback", on_date=today() + timedelta(days=1)
        )
    path.write_bytes(content("CZ").replace(b"Synthetic", b"Updated"))
    second = stage_directory(engine, "CZ", path, today(), today(), url)["version"]
    activate_directory(engine, "CZ", second, review_note="Second")
    activate_directory(engine, "CZ", first, review_note="Rollback")
    assert lookup_directory(engine, "CZ", "0100")[1]["name"] == "Synthetic Bank"
    with engine.connect() as conn:
        assert len(conn.execute(select(history)).all()) == 3
    path.write_bytes(content("CZ").replace(b"Synthetic", b"Future"))
    future = stage_directory(
        engine, "CZ", path, today() + timedelta(days=1), today() + timedelta(days=2), url
    )["version"]
    with pytest.raises(ValueError):
        activate_directory(engine, "CZ", future, review_note="Future")


def test_slovenian_longest_prefix_lookup(engine, tmp_path):
    path = tmp_path / "si.html"
    path.write_bytes(content("SI"))
    version = stage_directory(engine, "SI", path, today(), today(), SOURCES["SI"].page)["version"]
    activate_directory(engine, "SI", version, review_note="Reviewed")
    assert lookup_directory(engine, "SI", "01999")[1]["bank_code"] == "01"
    assert lookup_directory(engine, "SI", "91001")[1]["name"] == "Other Bank"
    assert lookup_directory(engine, "SI", "91002")[1] is None


@pytest.mark.parametrize(
    "url",
    [
        "http://www.cnb.cz/x",
        "https://evil.example/x",
        "https://www.cnb.cz.evil.example/x",
        "https://user@www.cnb.cz/x",
        "https://www.cnb.cz:444/x",
    ],
)
def test_official_url_allowlist(url):
    with pytest.raises(ValueError):
        validate_url("CZ", url)
    with pytest.raises(ValueError):
        OfficialRedirects("CZ").redirect_request(None, None, 302, "", {}, url)


def test_download_link_discovery_requires_unique_current_export():
    page = b'<a href="/uploads/FIK_KODAI_20261008-en.csv">csv</a>'
    assert discover_url("LT", page).endswith("FIK_KODAI_20261008-en.csv")
    with pytest.raises(ValueError):
        discover_url("LT", page + b'<a href="/FIK_KODAI_20261009-en.csv">csv</a>')
    page = (
        b'<a href="/bic_saraksts_2024_ENG.xls">hidden</a>'
        b'<a href="/bic_saraksts_2026_ENG1.xls">BIC list</a>'
    )
    assert discover_url("LV", page).endswith("2026_ENG1.xls")


def iban(country, bban):
    rearranged = bban + country + "00"
    digits = "".join(str(ord(c) - 55) if c.isalpha() else c for c in rearranged)
    return country + f"{98 - int(digits) % 97:02d}" + bban


EXAMPLES = {
    "NL": "TEST" + "0" * 10,
    "CH": "00100" + "0" * 12,
    "PL": "10100000" + "0" * 16,
    "LT": "10100" + "0" * 11,
    "BE": "000" + "0" * 9,
    "CZ": "0100" + "0" * 16,
    "LV": "TEST" + "0" * 13,
    "SI": "01999" + "0" * 10,
    "GR": "0110000" + "0" * 16,
}


@pytest.mark.parametrize("country", SOURCES)
def test_international_mod97_and_local_bank_api(engine, tmp_path, country):
    value = iban(country, EXAMPLES[country])
    assert validate_iban(value).iban_valid
    assert not validate_iban(value[:-1] + "1").iban_valid
    path = tmp_path / country
    path.write_bytes(content(country))
    version = stage_directory(
        engine, country, path, today(), today(), SOURCES[country].download or SOURCES[country].page
    )["version"]
    with TestClient(create_app(engine)) as client:
        response = client.post("/v1/validate", json={"iban": value}).json()
        assert response["iban_valid"] and response["bank_lookup_status"] == "unavailable"
        activate_directory(engine, country, version, review_note="Synthetic source test")
        response = client.post("/v1/validate", json={"iban": value}).json()
        assert response["bank_lookup_status"] == "found"
        assert response["bank"]["name"].startswith("Synthetic Bank")
        assert response["data"]["source"] == "Source: " + SOURCES[country].publisher
        assert len(client.get("/v1/countries").json()["countries"]) == len(COUNTRIES)


def test_provenance_dates_and_country_are_validated(engine, tmp_path):
    path = tmp_path / "ch.csv"
    path.write_bytes(content("CH"))
    url = SOURCES["CH"].download
    with pytest.raises(ValueError):
        stage_directory(engine, "CH", path, today() - timedelta(days=1), today(), url)
    first = stage_directory(engine, "CH", path, today(), today(), url)
    with pytest.raises(ValueError):
        stage_directory(engine, "CH", path, today(), today() + timedelta(days=1), url)
    with pytest.raises(ValueError):
        stage_directory(engine, "CH", path, today(), today(), SOURCES["CH"].page)
    with pytest.raises(ValueError):
        activate_directory(
            engine, "CH", first["version"], review_note="Reviewed", expected_previous_id="obsolete"
        )
    assert lookup_directory(engine, "CH", "00100")[0] is None


def test_deleted_swiss_iid_does_not_redirect_iban(engine, tmp_path):
    data = content("CH") + ("30100;" + str(today()) + ";Y;100;" + ";" * 16 + "\r\n").encode()
    path = tmp_path / "ch.csv"
    path.write_bytes(data)
    version = stage_directory(engine, "CH", path, today(), today(), SOURCES["CH"].download)[
        "version"
    ]
    activate_directory(engine, "CH", version, review_note="Reviewed")
    with TestClient(create_app(engine)) as client:
        result = client.post("/v1/validate", json={"iban": iban("CH", "30100" + "0" * 12)}).json()
        assert result["iban_valid"]
        assert result["bank_lookup_status"] == "deleted"
        assert result["bank_code_valid"] is False
        assert result["bank"] is None


def test_additive_schema_upgrade_keeps_german_data(engine, import_file):
    from openiban.directory_storage import active, records, versions
    from openiban.storage import initialize, metadata

    original = import_file(make_active=True)
    metadata.drop_all(engine, tables=[history, active, records, versions])
    initialize(engine)
    initialize(engine)
    assert lookup(engine, "12345678")[0]["version"] == original["version"]
    assert list_directory_versions(engine, "CZ") == []


def test_cli_stage_activate_and_list(engine, tmp_path, monkeypatch, capsys):
    from openiban.cli import main

    monkeypatch.setenv("OPENIBAN_DATABASE_URL", str(engine.url))
    path = tmp_path / "cz.csv"
    path.write_bytes(content("CZ"))
    monkeypatch.setattr(
        "sys.argv",
        [
            "openiban",
            "import-directory",
            "CZ",
            str(path),
            "--source-url",
            SOURCES["CZ"].download,
            "--valid-from",
            str(today()),
            "--valid-until",
            str(today()),
        ],
    )
    main()
    version = json.loads(capsys.readouterr().out)["version"]
    monkeypatch.setattr(
        "sys.argv",
        [
            "openiban",
            "activate-directory",
            "CZ",
            version,
            "--review-note",
            "Source, freshness and rights reviewed",
        ],
    )
    main()
    assert json.loads(capsys.readouterr().out)["status"] == "active"
    monkeypatch.setattr("sys.argv", ["openiban", "directory-versions", "CZ"])
    main()
    assert json.loads(capsys.readouterr().out)[0]["active"]


def test_bounded_download_and_http_failure(monkeypatch):
    from urllib.error import HTTPError

    from openiban.directory_download import fetch

    class Response(io.BytesIO):
        headers = {}

    class Opener:
        def open(self, request, timeout):
            assert timeout == 30
            return Response(b"x" * (MAX_BYTES + 1))

    monkeypatch.setattr("openiban.directory_download.build_opener", lambda handler: Opener())
    with pytest.raises(ValueError):
        fetch("CZ", SOURCES["CZ"].download)

    class FailingOpener:
        def open(self, request, timeout):
            raise HTTPError(request.full_url, 403, "Forbidden", {}, None)

    monkeypatch.setattr("openiban.directory_download.build_opener", lambda handler: FailingOpener())
    with pytest.raises(ValueError, match="unreachable"):
        fetch("CZ", SOURCES["CZ"].download)


def test_truncated_workbook_xml_fails_cleanly():
    data = xlsx([["Bank", "BIC", "Bank Identifier"], ["Synthetic Bank", "TESTGR22", "011"]])
    output = io.BytesIO()
    with ZipFile(io.BytesIO(data)) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            target.writestr(
                entry.filename,
                b"<broken"
                if entry.filename.endswith("sheet1.xml")
                else source.read(entry.filename),
            )
    with pytest.raises(ValueError):
        parse_directory("GR", output.getvalue())


@pytest.mark.parametrize(
    "row",
    [
        ["TESTNL22", "OTHER", "Synthetic Bank"],
        ["TESTBE22", "TEST", "Synthetic Bank"],
        ["TESTNL22", "TEST", ""],
        ["TESTNL22", "TEST", "=1+1"],
    ],
)
def test_nl_rejects_invalid_mappings_and_formulas(row):
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(content("NL")))
    for column, value in enumerate(row, 1):
        book.active.cell(3, column, value)
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(ValueError):
        parse_directory("NL", out.getvalue())


def test_nl_publication_date_duplicates_and_country_format():
    from openpyxl import load_workbook

    assert parse_directory("NL", content("NL")).published_on == today()
    assert not validate_iban(iban("NL", "1234" + "0" * 10)).iban_valid
    assert not validate_iban(iban("NL", "TEST" + "A" + "0" * 9)).iban_valid
    book = load_workbook(io.BytesIO(content("NL")))
    book.active.append(["TESTNL22", "TEST", "Duplicate Bank"])
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(ValueError):
        parse_directory("NL", out.getvalue())
