import io
from datetime import timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from openiban.api import create_app
from openiban.directories import parse_directory
from openiban.directory_download import discover_url
from openiban.directory_storage import checks, list_directory_versions, lookup_directory
from openiban.netherlands import update_nl
from openiban.storage import today

URL = "https://www.betaalvereniging.nl/wp-content/uploads/2025/11/BIC-lijst-NL.xlsx"


def workbook(count=60, published=None, name="Example"):
    book = Workbook()
    sheet = book.active
    published = published or today() - timedelta(days=1)
    sheet.append(
        [f"BIC-lijst-NL | BIC-list-NL (Laatste update | last update {published:%d-%m-%Y})"]
    )
    sheet.append(["BIC", "Identifier", "Betaaldienstverlener / Payment Service Provider"])
    for n in range(count):
        code = "AA" + chr(65 + n // 26) + chr(65 + n % 26)
        sheet.append([code + "NL2A", code, name + str(n)])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def test_discovery():
    assert discover_url("NL", f'<a href="{URL}">list</a>'.encode()) == URL
    for page in [
        b"<html/>",
        b'<a href="https://evil.test/BIC-lijst-NL.xlsx">x</a>',
        f'<a href="{URL}">x</a><a href="/other/BIC-lijst-NL.xlsx">y</a>'.encode(),
    ]:
        with pytest.raises(ValueError):
            discover_url("NL", page)


def test_update_and_freshness(engine, tmp_path, monkeypatch):
    monkeypatch.setattr("openiban.netherlands.download_directory", lambda _: (workbook(), URL))
    first = update_nl(engine, tmp_path)
    assert first["status"] == "updated"
    assert update_nl(engine, tmp_path)["status"] == "unchanged"
    assert len(list_directory_versions(engine, "NL")) == 1
    info, bank = lookup_directory(engine, "NL", "AAAA")
    assert bank["bic"] == "AAAANL2A"
    assert info["source"] == "Source: Betaalvereniging Nederland"
    assert "no guarantee" in info["disclaimer"]
    assert (
        lookup_directory(engine, "NL", "AAAA", today() + timedelta(days=8))[0]["status"]
        == "expired"
    )
    with TestClient(create_app(engine)) as client:
        # Known MOD-97 valid Dutch IBAN, absent from this synthetic snapshot.
        data = client.post("/v1/validate", json={"iban": "NL91ABNA0417164300"}).json()
        assert data["bank_lookup_status"] == "not_found"
        assert data["bank_code_valid"] is None
        assert data["data"]["checked_at"]
        assert data["data"]["disclaimer"]
    with engine.connect() as conn:
        checked = conn.execute(select(checks)).mappings().one()["checked_at"]
    for content in [
        b"broken",
        workbook(5),
        workbook(name="Changed"),
        workbook(published=today() - timedelta(days=2)),
        workbook(published=today() + timedelta(days=1)),
    ]:
        monkeypatch.setattr(
            "openiban.netherlands.download_directory", lambda _, c=content: (c, URL)
        )
        with pytest.raises(ValueError):
            update_nl(engine, tmp_path)
        assert lookup_directory(engine, "NL", "AAAA")[0]["version"] == first["version"]
        with engine.connect() as conn:
            assert conn.execute(select(checks.c.checked_at)).scalar_one() == checked


def test_small_update_and_removal(engine, tmp_path, monkeypatch):
    monkeypatch.setattr("openiban.netherlands.download_directory", lambda _: (workbook(), URL))
    update_nl(engine, tmp_path)
    monkeypatch.setattr("openiban.netherlands.download_directory", lambda _: (workbook(59), URL))
    assert update_nl(engine, tmp_path)["changed"] == 1
    assert len(list_directory_versions(engine, "NL")) == 2


def test_missing_bic_rejected():
    content = workbook()
    assert len(parse_directory("NL", content).records) == 60
    book = load_workbook(io.BytesIO(content))
    book.active["A3"] = "/"
    output = io.BytesIO()
    book.save(output)
    with pytest.raises(ValueError):
        parse_directory("NL", output.getvalue())


def test_activation_failure_rolls_back_freshness(engine, tmp_path, monkeypatch):
    monkeypatch.setattr("openiban.netherlands.download_directory", lambda _: (workbook(), URL))
    initial = update_nl(engine, tmp_path)
    monkeypatch.setattr("openiban.netherlands.download_directory", lambda _: (workbook(59), URL))

    def fail(*args, **kwargs):
        raise ValueError("simulated activation failure")

    monkeypatch.setattr("openiban.netherlands.activate_directory", fail)
    with pytest.raises(ValueError, match="simulated"):
        update_nl(engine, tmp_path)
    info, bank = lookup_directory(engine, "NL", "AAAA")
    assert info["version"] == initial["version"]
    assert info["checked_at"] == initial["checked_at"]
    assert bank


def test_overlapping_run_rejected(engine, tmp_path):
    from openiban.updater import locked

    with locked(tmp_path):
        with pytest.raises(ValueError, match="Another updater"):
            update_nl(engine, tmp_path)
